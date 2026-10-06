"""Word output — fills our own template (templates/test_protocol_template.docx) with the drafted rows.

Interface: build_docx(meta, records, out_path, template_path=None) -> dict(counts)
Requires python-docx (the only third-party dependency of TraceDraft). Values that are MISSING,
flagged by the verifier, or in conflict are written in red so a person cannot miss them.
The output is a DRAFT FOR REVIEW, never a released document.
"""
from __future__ import annotations

from pathlib import Path

COLUMNS = [("test_name", "Test Name"), ("req_id", "Req ID"), ("specification", "Specification"), ("n", "n"),
           ("acceptance", "Acceptance Criteria"), ("test_instruction", "Test Instruction")]
RED = (0xC0, 0x00, 0x00)


def build_docx(meta: dict, records: list[dict], out_path: str | Path, template_path: str | Path | None = None) -> dict:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document(str(template_path)) if template_path and Path(template_path).is_file() else Document()
    for p in doc.paragraphs:
        for k, v in {"{{PRODUCT}}": meta.get("product", ""), "{{CASE}}": meta.get("case_id", ""),
                     "{{RUN}}": meta.get("run_id", ""), "{{SNAPSHOT}}": meta.get("snapshot", "")[:16]}.items():
            if k in p.text:
                for r in p.runs:
                    r.text = r.text.replace(k, v)
    anchor = next((p for p in doc.paragraphs if "{{TABLES}}" in p.text), None)
    if anchor is not None:
        anchor.text = ""
    counts = {"cells": 0, "red": 0}
    by_cond: dict[str, list[dict]] = {}
    labels: dict[str, str] = {}
    for rec in records:
        by_cond.setdefault(rec["condition"], []).append(rec)
        labels[rec["condition"]] = rec["condition_label"]

    def add(par_or_cell_text_target, text, red=False):
        run = par_or_cell_text_target.add_run(text)
        run.font.size = Pt(9)
        if red:
            run.font.color.rgb = RGBColor(*RED)
            counts["red"] += 1
        return run

    for cond, recs in by_cond.items():
        heading = doc.add_heading(f"Test table — {labels[cond]}", level=2)
        if anchor is not None:
            anchor._p.addprevious(heading._p)
        tests: dict[str, dict] = {}
        for r in recs:
            tests.setdefault(r["test_name"], {})[r["field"]] = r
        table = doc.add_table(rows=1, cols=len(COLUMNS))
        table.style = "Table Grid"
        for i, (_, title) in enumerate(COLUMNS):
            add(table.rows[0].cells[i].paragraphs[0], title).bold = True
        for test_name, fields in tests.items():
            cells = table.add_row().cells
            add(cells[0].paragraphs[0], test_name)
            for i, (fname, _) in enumerate(COLUMNS[1:], 1):
                r = fields.get(fname)
                counts["cells"] += 1
                if r is None:
                    add(cells[i].paragraphs[0], "[MISSING]", red=True)
                    continue
                claim = r["claim"]
                flagged = r.get("risk", 0) >= 2
                if claim.get("value") in (None, ""):
                    add(cells[i].paragraphs[0], "[MISSING]", red=True)
                else:
                    add(cells[i].paragraphs[0], str(claim["value"]), red=flagged)
                if claim.get("evidence_id"):
                    add(cells[i].paragraphs[0], f"  [{claim['evidence_id']}]").font.size = Pt(7)
                if claim.get("conflict"):
                    add(cells[i].paragraphs[0], f"  CONFLICT: {claim['conflict']}", red=True)
        if anchor is not None:
            anchor._p.addprevious(table._tbl)
    doc.save(str(out_path))
    return counts
